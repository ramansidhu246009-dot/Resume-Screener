import io
import re
import pdfplumber
import docx

def normalize_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def extract_from_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(file_bytes)
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            text += textpage.get_text_bounded() + "\n"
    except Exception as e:
        import io
        import pdfplumber
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as fallback_e:
            raise ValueError(f"PDF extraction error: {e}, fallback: {fallback_e}")
    return normalize_text(text)

def extract_from_docx(file_bytes: bytes) -> str:
    text = ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        raise ValueError(f"DOCX extraction error: {e}")
    return normalize_text(text)

def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = filename.lower().split('.')[-1]
    if ext == 'pdf':
        return extract_from_pdf(file_bytes)
    elif ext == 'docx':
        return extract_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file type")
